from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path('/Users/mahartshorne/Documents/Codex/2026-09-15/referenced-chatgpt-conversation-this-is-an/outputs/carrier-configuration')
SOURCE = ROOT / 'docs/completion-rule-matrix.md'
OUT = ROOT / 'docs/AHM565_Configuration_Completion_Matrix.docx'

NAVY = '18345C'
BLUE = '2F63B8'
PALE_BLUE = 'EEF4FD'
PALE_ALT = 'F7F9FC'
BORDER = 'D9E1EC'
GREEN = '25633B'
GREEN_BG = 'E7F5EC'
AMBER = '745514'
AMBER_BG = 'FFF4DC'
RED = 'A12C2C'
RED_BG = 'FDEBEC'
GREY = '596579'
GREY_BG = 'EDF0F5'
WHITE = 'FFFFFF'
BLACK = '000000'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def borders(cell, color=BORDER, size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_el = tc_pr.find(qn('w:tcBorders'))
    if borders_el is None:
        borders_el = OxmlElement('w:tcBorders')
        tc_pr.append(borders_el)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders_el.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders_el.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_font(run, name='Aptos', size=8.2, bold=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text, size=8.2, color=BLACK):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        code = part.startswith('`') and part.endswith('`')
        clean = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(clean)
        set_font(run, 'Aptos Mono' if code else 'Aptos', size, bold, color)


def style_cell(cell, header=False, fill=None, align=None, size=7.8):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    borders(cell)
    if fill:
        shade(cell, fill)
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        if align is not None:
            p.alignment = align
        for r in p.runs:
            set_font(r, size=size, bold=header, color=WHITE if header else BLACK)


def add_table(doc, rows, widths=None, status_column=None, font_size=7.5, repeat_header=False):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if repeat_header:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = value
        style_cell(cell, header=True, fill=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, size=7.6)
        if widths:
            cell.width = Inches(widths[i])
    for ridx, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            add_inline(p, value, font_size)
            fill = PALE_ALT if ridx % 2 == 0 else WHITE
            if status_column == i:
                upper = value.upper()
                if 'DONE' == upper or 'CONFIGURED' == upper:
                    fill = GREEN_BG
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(GREEN)
                        run.font.bold = True
                elif 'PARTIAL' in upper or 'PENDING' in upper:
                    fill = AMBER_BG
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(AMBER)
                        run.font.bold = True
                elif 'SKIPPED' in upper or 'SUPPORTED' in upper or 'OUT OF SCOPE' in upper:
                    fill = GREY_BG
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(GREY)
                        run.font.bold = True
                elif 'NOT STARTED' in upper or 'INCOMPLETE' in upper:
                    fill = RED_BG
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(RED)
                        run.font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_cell(cells[i], fill=fill, size=font_size)
            if widths:
                cells[i].width = Inches(widths[i])
    table.rows[0].height = Cm(0.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level, page_break=False):
    p = doc.add_paragraph(style=f'Heading {level}')
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=15 if level == 1 else 11, bold=True, color=BLACK)


def normalize(text):
    return (text.replace('—', '-').replace('–', '-').replace('‑', '-').replace('“', '"').replace('”', '"').replace('’', "'"))


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.35)
section.bottom_margin = Cm(1.3)
section.left_margin = Cm(1.4)
section.right_margin = Cm(1.4)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(9)
styles['Normal'].font.color.rgb = RGBColor.from_string(BLACK)
for s in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
    styles[s].font.name = 'Aptos Display'
    styles[s]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    styles[s]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    styles[s].font.color.rgb = RGBColor.from_string(BLACK)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(5)
r = title.add_run('AHM565 Configuration Completion Rule Matrix')
set_font(r, 'Aptos Display', 25, True, BLACK)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(10)
add_inline(sub, 'Implemented B and C screens, completion rules and remaining status work  |  19 September 2026', 10, GREY)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(10)
add_inline(intro, 'Purpose. This document records the agreed Red, Amber and Green completion rules and shows which operational screens and database workflows have been delivered. A screen marked DONE is available for use; the Status Rollout column separately records whether its automatic completion evaluator and badge are already live.', 9)

add_heading(doc, 'Implementation Progress', 1)
progress_rows = [
    ['Page or section', 'Operational screen', 'Database workflow', 'Completion rules', 'Status rollout', 'Remaining work'],
    ['B1 Standard Units and Codes', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add review metadata where empty density values can be deliberate; implement evaluator and badges.'],
    ['B2 Crew and Crew Baggage Weights', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Implement section/page evaluator and badges.'],
    ['B3 Passenger and Hand Baggage Weights', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add review metadata for optional variations and class weights; implement evaluator and badges.'],
    ['B4 Standard Baggage Weights and Planning', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add review metadata for optional weight sets and planning assumptions; implement evaluator and badges.'],
    ['B5 ULD Specifications', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add reviewed/applicability state; implement evaluator and badges.'],
    ['C1 Aircraft Type or Fleet', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Replace the current simple configured indicator with the shared rule evaluator and badge.'],
    ['C2 and C3 Loadsheet Output', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add review metadata for Lower Loadsheet Information; implement evaluator and badges.'],
    ['C4 Basic Index and MAC Formula', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Implement evaluator and badge.'],
    ['C5.1 Balance Envelope', 'DONE', 'DONE', 'DONE', 'DONE', 'No current rollout work. Revalidation includes maxima, MAX rows and the fleet-DOW lower-bound coverage rule.'],
    ['C6 Curtailments', 'PLACEHOLDER', 'SKIPPED', 'SKIPPED', 'SKIPPED', 'Activate only when C6 is brought into scope. Navigation currently proceeds from C5.1 to C7.'],
    ['C7 Ideal Trim', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add reviewed/applicability state; implement evaluator and badge.'],
    ['C7 Tipping Limits', 'DONE', 'DONE', 'DEFINED', 'PENDING', 'Add reviewed/applicability state; implement evaluator and badge.'],
    ['C7 Lateral Imbalance', 'PLACEHOLDER', 'UNSUPPORTED', 'EXCLUDED', 'UNSUPPORTED', 'Future scope; excluded from completion aggregation.'],
    ['Aircraft Type aggregate', 'NOT STARTED', 'NOT STARTED', 'DEFINED', 'PENDING', 'Aggregate active supported C pages after page evaluators are live.'],
    ['Carrier aggregate', 'NOT STARTED', 'NOT STARTED', 'DEFINED', 'PENDING', 'Aggregate B pages and aircraft after page evaluators are live.'],
]
add_table(doc, progress_rows, widths=[1.72, 1.05, 1.1, 1.0, 1.25, 4.0], status_column=4, font_size=7.4, repeat_header=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
add_inline(p, '**Current position:** all operational B1-B5 and C1-C5.1/C7 workflows covered by this matrix are built. The shared badge and full automatic evaluator are live on C5.1. C6 is deliberately skipped, and C7 Lateral Imbalance is unsupported.', 9)

add_heading(doc, 'Status Key', 1)
key_rows = [
    ['Colour', 'Display text', 'Meaning'],
    ['Red', 'INCOMPLETE', 'A mandatory prerequisite is missing, the section has not been reviewed, or no meaningful configuration has been saved.'],
    ['Amber', 'PARTIALLY CONFIGURED', 'Some relevant data has been saved, but at least one applicable completion rule is not satisfied.'],
    ['Green', 'CONFIGURED', 'Every applicable completion rule is satisfied by valid saved data.'],
    ['Neutral grey', 'UNSUPPORTED / OUT OF SCOPE / SKIPPED', 'Excluded from completion aggregation.'],
]
table = add_table(doc, key_rows, widths=[1.2, 2.65, 6.25], font_size=8)
for row in table.rows[1:]:
    txt = row.cells[1].text
    if txt == 'INCOMPLETE': shade(row.cells[1], RED_BG)
    elif txt == 'PARTIALLY CONFIGURED': shade(row.cells[1], AMBER_BG)
    elif txt == 'CONFIGURED': shade(row.cells[1], GREEN_BG)
    else: shade(row.cells[1], GREY_BG)

add_heading(doc, 'Detailed Completion Rules', 1)

text = normalize(SOURCE.read_text())
# The document title and initial status summary are already supplied above.
lines = text.splitlines()
i = 1
while i < len(lines):
    line = lines[i].rstrip()
    if line.startswith('# '):
        i += 1
        continue
    if line.startswith('## '):
        heading = line[3:].strip()
        is_page_section = bool(re.match(r'^[BC]\d', heading))
        if is_page_section:
            doc.add_page_break()
        add_heading(doc, heading, 1)
        if is_page_section:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(1)
        i += 1
        continue
    if line.startswith('### '):
        add_heading(doc, line[4:].strip(), 2)
        i += 1
        continue
    if line == '---':
        i += 1
        continue
    if line.startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].startswith('|'):
            table_lines.append(lines[i])
            i += 1
        parsed = []
        for tl in table_lines:
            vals = [normalize(v.strip()) for v in tl.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-{3,}:?', v or '') for v in vals):
                continue
            parsed.append(vals)
        if parsed:
            count = len(parsed[0])
            if count == 5:
                widths = [1.75, 1.32, 2.18, 2.35, 2.45]
            elif count == 4:
                widths = [1.8, 2.45, 2.8, 3.0]
            elif count == 3:
                widths = [1.5, 2.4, 6.0]
            else:
                widths = [10.0 / count] * count
            add_table(doc, parsed, widths=widths, font_size=7.15 if count >= 5 else 7.7)
        continue
    if re.match(r'^\d+\. ', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, line, 8.4)
        i += 1
        continue
    if line.startswith('- '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, '•\u2002' + line[2:], 8.4)
        i += 1
        continue
    if line:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        add_inline(p, line, 8.5)
    i += 1

# Add the corrected DOW coverage rule as an explicit implementation note.
add_heading(doc, 'C5.1 Effective DOW Coverage Rule', 1)
p = doc.add_paragraph()
add_inline(p, 'For every Forward and Aft envelope, the first or lowest Weight must be less than or equal to the smallest available Dry Operating Weight for the matching carrier, Aircraft Type and Series/Sub-Type. If no fleet DOW is available, Standard Fleet Weight is used. This allows the envelope to contain weights below DOW while ensuring that an empty ferry flight, where Actual Zero Fuel Weight equals DOW, is covered by the published envelope.', 8.7)

# Footer with title and dynamic page number.
footer = section.footer
tbl = footer.add_table(rows=1, cols=2, width=Inches(10.0))
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
left, right = tbl.rows[0].cells
left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
lr = left.paragraphs[0].add_run('AHM565 Configuration Completion Rule Matrix')
set_font(lr, size=7.5, color=GREY)
rr = right.paragraphs[0].add_run('Page ')
set_font(rr, size=7.5, color=GREY)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
right.paragraphs[0]._p.append(fld)

doc.core_properties.title = 'AHM565 Configuration Completion Rule Matrix'
doc.core_properties.subject = 'Implementation progress and completion rules for AHM565 B and C configuration screens'
doc.core_properties.author = 'Carrier Configuration Project'
doc.save(OUT)
print(OUT)
